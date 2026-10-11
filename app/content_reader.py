"""Extract text from local PDF, Word (.docx) and Excel (.xlsx) documents.

The heavy readers (pypdf, python-docx, openpyxl) are imported lazily so Jarvis
starts and runs normally when they are not installed yet; the resulting tool
error tells the user to install the project requirements.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

MAX_TEXT_CHARACTERS = 50_000
MAX_PDF_PAGES = 200
MAX_XLSX_ROWS = 10_000


def read_document(path: str) -> dict[str, Any]:
    """Return the text content of a local document as tool output."""
    target = Path(path).expanduser()
    if not target.is_file():
        return {"success": False, "error": f"File not found: {target}"}
    suffix = target.suffix.lower()
    try:
        if suffix == ".pdf":
            text = _read_pdf(target)
        elif suffix == ".docx":
            text = _read_docx(target)
        elif suffix in {".xlsx", ".xlsm"}:
            text = _read_xlsx(target)
        else:
            text = _read_text(target)
    except ImportError as exc:
        return {"success": False, "error": (
            f"Reading {suffix or 'this'} files needs an extra library ({exc}). "
            "Install the project requirements with: pip install -r requirements.txt")}
    except Exception as exc:
        return {"success": False, "error": f"Could not read {target.name}: {exc}"}
    truncated = len(text) > MAX_TEXT_CHARACTERS
    return {
        "success": True,
        "path": str(target.resolve()),
        "characters": len(text),
        "truncated": truncated,
        "text": text[:MAX_TEXT_CHARACTERS],
    }


def _read_text(path: Path) -> str:
    # Reject binary content early so unsupported formats return a clear error
    # instead of a page of replacement characters.
    if b"\x00" in path.read_bytes()[:8192]:
        raise ValueError(f"{path.suffix or 'binary'} files are not supported; use .txt, .md, .pdf, .docx, or .xlsx")
    return path.read_text(encoding="utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = []
    for index, page in enumerate(reader.pages):
        if index >= MAX_PDF_PAGES:
            pages.append("... (more pages omitted)")
            break
        pages.append(page.extract_text() or "")
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"== Sheet: {sheet.title} ==")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_XLSX_ROWS:
                    parts.append("... (more rows omitted)")
                    break
                values = [str(value) for value in row if value is not None and str(value).strip()]
                if values:
                    parts.append(", ".join(values))
    finally:
        workbook.close()
    return "\n".join(parts)