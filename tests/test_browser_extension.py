from __future__ import annotations

import pytest

from app.computer.browser import BrowserController
from app.content_reader import read_document
from app.permissions.policy import PermissionPolicy
from app.tooling import ToolRegistry


class FakeResponse:
    def __init__(self, status: int | None = 200) -> None:
        self.status = status


class FakeLocator:
    def count(self) -> int:
        return 1

    def click(self, **_kwargs) -> None:
        pass

    def fill(self, _text, **_kwargs) -> None:
        pass

    def focus(self, **_kwargs) -> None:
        pass

    def scroll_into_view_if_needed(self, **_kwargs) -> None:
        pass

    def inner_text(self, **_kwargs) -> str:
        return "Example page body"


class FakeKeyboard:
    def __init__(self, page) -> None:
        self.page = page

    def press(self, key: str) -> None:
        self.page.pressed_keys.append(key)


class FakeMouse:
    def __init__(self, page) -> None:
        self.page = page

    def wheel(self, delta_x: int, delta_y: int) -> None:
        self.page.wheel_moves.append((delta_x, delta_y))


class FakePage:
    def __init__(self) -> None:
        self.url = "https://example.com"
        self.history = ["https://example.com", "https://example.org"]
        self.history_index = 0
        self.pressed_keys: list[str] = []
        self.wheel_moves: list[tuple[int, int]] = []
        self.load_states: list[str] = []
        self.waited_ms: list[int] = []
        self.keyboard = FakeKeyboard(self)
        self.mouse = FakeMouse(self)

    def set_default_timeout(self, _milliseconds: int) -> None:
        pass

    def title(self) -> str:
        return "Example"

    def goto(self, url: str, **_kwargs) -> FakeResponse:
        self.url = url
        return FakeResponse(200)

    def locator(self, _selector: str) -> FakeLocator:
        return FakeLocator()

    def go_back(self, **_kwargs) -> FakeResponse | None:
        if self.history_index > 0:
            self.history_index -= 1
            self.url = self.history[self.history_index]
            return FakeResponse(200)
        return None

    def go_forward(self, **_kwargs) -> FakeResponse | None:
        if self.history_index < len(self.history) - 1:
            self.history_index += 1
            self.url = self.history[self.history_index]
            return FakeResponse(200)
        return None

    def wait_for_load_state(self, *, state: str, timeout: int) -> None:
        self.load_states.append(state)

    def wait_for_timeout(self, milliseconds: int) -> None:
        self.waited_ms.append(milliseconds)

    def eval_on_selector_all(self, selector: str, _expression: str, arg: int):
        assert selector == "a[href]"
        raw = [
            ("Home", "https://example.com/"),
            ("About", "https://example.com/about"),
            ("", "https://example.com/contact"),
            ("Mail", "mailto:test@example.com"),
            ("Home", "https://example.com/"),
        ]
        result = []
        for text, href in raw[:arg]:
            if href.startswith("http://") or href.startswith("https://"):
                result.append({"text": text, "href": href})
        return result


class FakeBrowser:
    def __init__(self) -> None:
        self.page = FakePage()
        self.closed = False

    def new_page(self) -> FakePage:
        return self.page

    def close(self) -> None:
        self.closed = True


class FakePlaywright:
    def __init__(self) -> None:
        self.chromium = self
        self.browser = FakeBrowser()
        self.launches: list[dict] = []

    def start(self):
        return self

    def launch(self, **options):
        self.launches.append(options)
        return self.browser

    def stop(self) -> None:
        pass


@pytest.fixture()
def browser_controller():
    playwright = FakePlaywright()
    controller = BrowserController(playwright_factory=lambda: playwright)
    opened = controller.open_browser("https://example.com")
    assert opened["success"]
    return controller, playwright


def test_browser_go_back_and_forward(browser_controller):
    controller, playwright = browser_controller
    page = playwright.browser.page
    page.history_index = 1
    page.url = "https://example.org"

    gone_back = controller.go_back()
    assert gone_back["success"]
    assert gone_back["url"] == "https://example.com"
    assert gone_back["status"] == 200

    gone_forward = controller.go_forward()
    assert gone_forward["success"]
    assert gone_forward["url"] == "https://example.org"


def test_press_browser_key_focuses_then_presses(browser_controller):
    controller, playwright = browser_controller
    page = playwright.browser.page

    result = controller.press_browser_key("Enter", selector="input[name=q]")
    assert result["success"]
    assert result["key"] == "Enter"
    assert page.pressed_keys == ["Enter"]

    result = controller.press_browser_key("Tab")
    assert result["success"]
    assert page.pressed_keys == ["Enter", "Tab"]


def test_press_browser_key_rejects_unknown_keys(browser_controller):
    controller, _ = browser_controller
    result = controller.press_browser_key("Ctrl+Shift+Delete")
    assert result["success"] is False
    assert "must be one of" in result["error"]


def test_browser_scroll_wheel_and_into_view(browser_controller):
    controller, playwright = browser_controller
    page = playwright.browser.page

    wheel = controller.scroll_browser(delta_x=0, delta_y=500)
    assert wheel["success"] and wheel["mode"] == "wheel"
    assert page.wheel_moves == [(0, 500)]

    into_view = controller.scroll_browser(selector="input[name=q]")
    assert into_view["success"] and into_view["mode"] == "into_view"

    oversized = controller.scroll_browser(delta_y=50_000)
    assert oversized["success"] is False
    assert "limited" in oversized["error"]


def test_browser_wait_states_and_delays(browser_controller):
    controller, playwright = browser_controller
    page = playwright.browser.page

    loaded = controller.wait_browser(state="networkidle")
    assert loaded["success"] and loaded["waited_for"] == "networkidle"
    assert page.load_states == ["networkidle"]

    paused = controller.wait_browser(milliseconds=1200)
    assert paused["success"] and paused["waited_ms"] == 1200
    assert page.waited_ms == [1200]

    bad = controller.wait_browser(state="frozen")
    assert bad["success"] is False

    too_long = controller.wait_browser(milliseconds=90_000)
    assert too_long["success"] is False


def test_extract_links_deduplicates(browser_controller):
    controller, _ = browser_controller
    result = controller.extract_links(limit=50)
    assert result["success"]
    assert result["count"] == 3
    hrefs = [link["href"] for link in result["links"]]
    assert hrefs == [
        "https://example.com/",
        "https://example.com/about",
        "https://example.com/contact",
    ]


def test_browser_wait_requires_an_open_page():
    class NeverStarted:
        def __call__(self):
            raise AssertionError("Must not launch a browser for a pure wait.")

    controller = BrowserController(playwright_factory=NeverStarted())
    result = controller.wait_browser(milliseconds=100)
    assert result["success"] is False
    assert "Open a browser" in result["error"]


def test_read_document_text_file(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Hello Jarvis\nsecond line", encoding="utf-8")
    result = read_document(str(path))
    assert result["success"] is True
    assert "Hello Jarvis" in result["text"]
    assert result["characters"] == len(result["text"])


def test_read_document_rejects_binary_without_garbage(tmp_path):
    path = tmp_path / "blob.dat"
    path.write_bytes(b"\x00\x01\x02binary")
    result = read_document(str(path))
    assert result["success"] is False
    assert "not supported" in result["error"]


def test_read_document_missing_file():
    result = read_document("C:/definitely/not/here.pdf")
    assert result["success"] is False
    assert "File not found" in result["error"]


def test_read_document_pdf(tmp_path):
    pytest.importorskip("pypdf")

    def build_pdf(text: str) -> bytes:
        escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT /F1 12 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]
        out = bytearray(b"%PDF-1.4\n")
        offsets: list[int] = []
        for index, body in enumerate(objects, start=1):
            offsets.append(len(out))
            out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
        xref = len(out)
        out += f"xref\n0 {len(objects) + 1}\n".encode() + b"0000000000 65535 f \n"
        for offset in offsets:
            out += f"{offset:010d} 00000 n \n".encode()
        out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
        return bytes(out)

    path = tmp_path / "sample.pdf"
    path.write_bytes(build_pdf("Hello PDF world"))
    result = read_document(str(path))
    assert result["success"] is True, result
    assert "Hello PDF world" in result["text"]


def test_read_document_docx(tmp_path):
    pytest.importorskip("docx")
    from docx import Document

    path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Hello Word document")
    document.save(str(path))

    result = read_document(str(path))
    assert result["success"] is True, result
    assert "Hello Word document" in result["text"]


def test_read_document_xlsx(tmp_path):
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    path = tmp_path / "sample.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "Name"
    sheet["B1"] = "Age"
    sheet["A2"] = "Ada"
    sheet["B2"] = 36
    workbook.save(str(path))

    result = read_document(str(path))
    assert result["success"] is True, result
    assert "== Sheet: Data ==" in result["text"]
    assert "Name, Age" in result["text"]
    assert "Ada, 36" in result["text"]


def test_new_browser_and_document_tools_registered_with_levels():
    registry = ToolRegistry()
    policy = PermissionPolicy()

    expected = {
        "browser_go_back": "moderate",
        "browser_go_forward": "moderate",
        "press_browser_key": "dangerous",
        "browser_scroll": "moderate",
        "browser_wait": "safe",
        "extract_links": "moderate",
        "read_document": "safe",
    }
    for name, level in expected.items():
        assert name in registry.tools, name
        assert registry.tools[name].permission_level == level, name
        assert policy.confirm_required(name, level) == (level != "safe")
        assert policy.allows(name, level, confirm=True) is True